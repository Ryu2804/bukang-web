import hashlib
import io
import random
import ssl
from collections.abc import Callable, Sequence
from concurrent.futures import ThreadPoolExecutor
from contextlib import ExitStack
from pathlib import Path
from tempfile import NamedTemporaryFile, TemporaryDirectory, TemporaryFile, gettempdir
from threading import Lock
from typing import BinaryIO
from urllib.parse import urlparse
from urllib.request import urlopen

import certifi
import httpx
from PIL import Image, ImageOps

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

from app.infrastructure.persistence.models.student import StudentModel


PDF_WIDTH = 1860.0
PDF_HEIGHT = 2631.0
SLIDE_WIDTH_INCHES = 10.0
SLIDE_HEIGHT_INCHES = SLIDE_WIDTH_INCHES * (PDF_HEIGHT / PDF_WIDTH)
SCALE_X = SLIDE_WIDTH_INCHES / PDF_WIDTH
SCALE_Y = SLIDE_HEIGHT_INCHES / PDF_HEIGHT
FOTO_PADDING = 12
LABEL_VALUE_GAP = 58
BACKGROUND_PATH = Path(__file__).with_name("page3_clean_bg.png")
FONT_NAME = "Garet"
FONT_SIZE_LABEL = Pt(24)
FONT_SIZE_VALUE = Pt(18)
COLOR_LABEL = RGBColor(0x00, 0x00, 0x00)
COLOR_VALUE = RGBColor(0x33, 0x33, 0x33)
COLOR_PHOTO_BORDER = RGBColor(0x00, 0x9E, 0x60)
FIRST_IMPRESSIONS = ["keren", "baik", "jago", "ramah"]
PHOTO_MAX_EDGE = 1600
PHOTO_JPEG_QUALITY = 85
PHOTO_DOWNLOAD_WORKERS = 16
PHOTO_DOWNLOAD_CHUNK_SIZE = 64 * 1024
PHOTO_CACHE_DIR = (
    Path(gettempdir())
    / f"bukang-pptx-photo-cache-{PHOTO_MAX_EDGE}-q{PHOTO_JPEG_QUALITY}"
)
_SSL_CONTEXT = None
_SSL_CONTEXT_LOCK = Lock()

SLOT_TOP = {
    "foto_outer": (296, 287, 1267, 493),
    "label_nama": (207, 812),
    "label_nrp": (207, 923),
    "label_prodi": (207, 1035),
    "label_asal": (207, 1147),
    "label_hobi": (207, 1259),
    "label_fi": (1101, 812),
    "nama": (207, 812 + LABEL_VALUE_GAP, 850),
    "nrp": (207, 923 + LABEL_VALUE_GAP, 850),
    "prodi": (207, 1035 + LABEL_VALUE_GAP, 850),
    "asal_daerah": (207, 1147 + LABEL_VALUE_GAP, 850),
    "hobi": (207, 1259 + LABEL_VALUE_GAP, 850),
    "first_impression": (1101, 812 + LABEL_VALUE_GAP, 560),
}

SLOT_BOTTOM = {
    "foto_outer": (296, 1393, 1267, 493),
    "label_nama": (207, 1918),
    "label_nrp": (207, 2030),
    "label_prodi": (207, 2142),
    "label_asal": (207, 2254),
    "label_hobi": (207, 2366),
    "label_fi": (1101, 1918),
    "nama": (207, 1918 + LABEL_VALUE_GAP, 850),
    "nrp": (207, 2030 + LABEL_VALUE_GAP, 850),
    "prodi": (207, 2142 + LABEL_VALUE_GAP, 850),
    "asal_daerah": (207, 2254 + LABEL_VALUE_GAP, 850),
    "hobi": (207, 2366 + LABEL_VALUE_GAP, 850),
    "first_impression": (1101, 1918 + LABEL_VALUE_GAP, 560),
}


def pdf_to_inches(x, y, width=None, height=None):
    left = Inches(x * SCALE_X)
    top = Inches(y * SCALE_Y)
    if width is not None and height is not None:
        return left, top, Inches(width * SCALE_X), Inches(height * SCALE_Y)
    return left, top


def get_ssl_context():
    global _SSL_CONTEXT
    if _SSL_CONTEXT is None:
        with _SSL_CONTEXT_LOCK:
            if _SSL_CONTEXT is None:
                _SSL_CONTEXT = ssl.create_default_context(cafile=certifi.where())
    return _SSL_CONTEXT


def load_photo(photo_url: str | None, *, opener: Callable = urlopen) -> bytes | None:
    if not photo_url:
        return None

    parsed = urlparse(photo_url)
    if not parsed.scheme and parsed.path.startswith("/uploads/"):
        local_path = Path("uploads") / Path(parsed.path).name
        try:
            return local_path.read_bytes()
        except OSError:
            return None

    is_hugging_face_bucket = (
        parsed.scheme == "https"
        and parsed.hostname == "huggingface.co"
        and parsed.path.startswith("/buckets/")
    )
    if is_hugging_face_bucket:
        try:
            with opener(photo_url, timeout=15, context=get_ssl_context()) as response:
                return response.read()
        except (OSError, ValueError):
            return None
    return None


def optimize_photo(photo_bytes: bytes | None) -> bytes | None:
    if not photo_bytes:
        return None

    try:
        with Image.open(io.BytesIO(photo_bytes)) as source:
            if source.format == "JPEG":
                source.draft("RGB", (PHOTO_MAX_EDGE, PHOTO_MAX_EDGE))
            orientation = source.getexif().get(274, 1)
            oriented = (
                ImageOps.exif_transpose(source)
                if orientation in {2, 3, 4, 5, 6, 7, 8}
                else source
            )
            oriented.thumbnail(
                (PHOTO_MAX_EDGE, PHOTO_MAX_EDGE),
                Image.Resampling.LANCZOS,
            )
            if oriented.mode in ("RGBA", "LA"):
                rgba = oriented.convert("RGBA")
                image = Image.new("RGB", rgba.size, "white")
                image.paste(rgba, mask=rgba.getchannel("A"))
            elif oriented.mode == "RGB":
                image = oriented
            else:
                image = oriented.convert("RGB")

            output = io.BytesIO()
            image.save(
                output,
                format="JPEG",
                quality=PHOTO_JPEG_QUALITY,
                optimize=True,
            )
            return output.getvalue()
    except (OSError, ValueError):
        return None


def spool_photo(
    index: int,
    photo_url: str | None,
    directory: Path,
    photo_loader: Callable[[str | None], bytes | None],
) -> Path | None:
    photo_bytes = photo_loader(photo_url)
    if not photo_bytes:
        return None

    path = directory / f"{index:06d}.photo"
    try:
        path.write_bytes(photo_bytes)
    except OSError:
        return None
    return path


def create_photo_http_client() -> httpx.Client:
    return httpx.Client(
        follow_redirects=True,
        limits=httpx.Limits(
            max_connections=PHOTO_DOWNLOAD_WORKERS,
            max_keepalive_connections=PHOTO_DOWNLOAD_WORKERS,
        ),
        timeout=httpx.Timeout(60.0, connect=15.0),
        verify=get_ssl_context(),
    )


def stream_photo_to_disk(
    index: int,
    photo_url: str | None,
    directory: Path,
    client: httpx.Client,
) -> Path | None:
    if not photo_url:
        return None

    parsed = urlparse(photo_url)
    local_path = None
    if not parsed.scheme and parsed.path.startswith("/uploads/"):
        local_path = Path("uploads") / Path(parsed.path).name
    is_hugging_face_bucket = (
        parsed.scheme == "https"
        and parsed.hostname == "huggingface.co"
        and parsed.path.startswith("/buckets/")
    )
    if local_path is None and not is_hugging_face_bucket:
        return None

    path = directory / f"{index:06d}.photo"
    try:
        with path.open("wb") as destination:
            if local_path is not None:
                with local_path.open("rb") as source:
                    while chunk := source.read(PHOTO_DOWNLOAD_CHUNK_SIZE):
                        destination.write(chunk)
            else:
                with client.stream("GET", photo_url) as response:
                    response.raise_for_status()
                    for chunk in response.iter_bytes(PHOTO_DOWNLOAD_CHUNK_SIZE):
                        destination.write(chunk)
        return path if path.stat().st_size > 0 else None
    except (httpx.HTTPError, OSError, ValueError):
        path.unlink(missing_ok=True)
        return None


def consume_spooled_photo(path: Path | None) -> bytes | None:
    if path is None:
        return None
    try:
        return path.read_bytes()
    except OSError:
        return None
    finally:
        path.unlink(missing_ok=True)


def cached_photo_path(photo_url: str | None, directory: Path) -> Path | None:
    if not photo_url:
        return None
    key = hashlib.sha256(photo_url.encode("utf-8")).hexdigest()
    path = directory / f"{key}.jpg"
    try:
        return path if path.is_file() and path.stat().st_size > 0 else None
    except OSError:
        return None


def store_cached_photo(photo_url: str | None, photo_bytes: bytes, directory: Path) -> None:
    if not photo_url:
        return
    directory.mkdir(mode=0o700, parents=True, exist_ok=True)
    destination = directory / f"{hashlib.sha256(photo_url.encode('utf-8')).hexdigest()}.jpg"
    temporary_path = None
    try:
        with NamedTemporaryFile(mode="w+b", dir=directory, delete=False) as temporary:
            temporary.write(photo_bytes)
            temporary_path = Path(temporary.name)
        temporary_path.replace(destination)
    except OSError:
        if temporary_path is not None:
            temporary_path.unlink(missing_ok=True)


def add_text(
    slide,
    text,
    pdf_x,
    pdf_y,
    pdf_max_width,
    font_size=FONT_SIZE_VALUE,
    bold=False,
    color=COLOR_VALUE,
):
    left, top = pdf_to_inches(pdf_x, pdf_y)
    width = Inches(pdf_max_width * SCALE_X)
    height = Inches(55 * SCALE_Y)
    text_box = slide.shapes.add_textbox(left, top, width, height)
    text_frame = text_box.text_frame
    text_frame.word_wrap = True
    paragraph = text_frame.paragraphs[0]
    paragraph.text = text
    paragraph.font.name = FONT_NAME
    paragraph.font.size = font_size
    paragraph.font.bold = bold
    paragraph.font.color.rgb = color


def add_empty_foto_frame(slide, slot):
    x, y, width, height = slot["foto_outer"]
    left, top, outer_width, outer_height = pdf_to_inches(x, y, width, height)
    border = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, outer_width, outer_height)
    border.fill.solid()
    border.fill.fore_color.rgb = COLOR_PHOTO_BORDER
    border.line.fill.background()

    inner = pdf_to_inches(
        x + FOTO_PADDING,
        y + FOTO_PADDING,
        width - (2 * FOTO_PADDING),
        height - (2 * FOTO_PADDING),
    )
    placeholder = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, *inner)
    placeholder.fill.solid()
    placeholder.fill.fore_color.rgb = RGBColor(0xF0, 0xF0, 0xF0)
    placeholder.line.fill.background()


def add_foto_with_border(slide, foto_bytes, slot):
    try:
        with Image.open(io.BytesIO(foto_bytes)) as image:
            image_width, image_height = image.size
    except Exception:
        add_empty_foto_frame(slide, slot)
        return

    x, y, width, height = slot["foto_outer"]
    max_width = width - (2 * FOTO_PADDING)
    max_height = height - (2 * FOTO_PADDING)
    scale = min(max_width / image_width, max_height / image_height)
    fitted_width = image_width * scale
    fitted_height = image_height * scale
    image_x = x + ((width - fitted_width) / 2)
    image_y = y + ((height - fitted_height) / 2)

    border = pdf_to_inches(
        image_x - FOTO_PADDING,
        image_y - FOTO_PADDING,
        fitted_width + (2 * FOTO_PADDING),
        fitted_height + (2 * FOTO_PADDING),
    )
    border_shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, *border)
    border_shape.fill.solid()
    border_shape.fill.fore_color.rgb = COLOR_PHOTO_BORDER
    border_shape.line.fill.background()

    picture_box = pdf_to_inches(image_x, image_y, fitted_width, fitted_height)
    slide.shapes.add_picture(io.BytesIO(foto_bytes), *picture_box)


def fill_slot(slide, slot, student, foto_bytes=None):
    nama = student.get("nama") or "-"
    nrp = student.get("nrp") or "-"
    prodi = student.get("prodi") or "-"
    asal = student.get("asal") or "-"
    hobi = student.get("hobi") or "-"
    impression = random.choice(FIRST_IMPRESSIONS)

    if foto_bytes:
        add_foto_with_border(slide, foto_bytes, slot)
    else:
        add_empty_foto_frame(slide, slot)

    labels = (
        ("label_nama", "Nama Lengkap", 600),
        ("label_nrp", "NRP", 600),
        ("label_prodi", "Prodi", 600),
        ("label_asal", "Asal Daerah", 600),
        ("label_hobi", "Hobi", 600),
        ("label_fi", "First Impression", 560),
    )
    for key, text, max_width in labels:
        x, y = slot[key]
        add_text(
            slide,
            text,
            x,
            y,
            max_width,
            font_size=FONT_SIZE_LABEL,
            bold=True,
            color=COLOR_LABEL,
        )

    values = (
        ("nama", nama),
        ("nrp", nrp),
        ("prodi", prodi),
        ("asal_daerah", asal),
        ("hobi", hobi),
        ("first_impression", impression),
    )
    for key, value in values:
        x, y, max_width = slot[key]
        add_text(slide, value, x, y, max_width)


def _student_record(student: StudentModel) -> dict[str, str]:
    return {
        "nama": student.name or "",
        "nrp": student.nrp or "",
        "prodi": student.major or "",
        "asal": student.hometown or "",
        "hobi": student.hobbies or "",
        "first_impression": student.first_impression or "",
    }


def generate_pptx(
    students: Sequence[StudentModel],
    *,
    photo_loader: Callable[[str | None], bytes | None] | None = None,
    background_path: Path = BACKGROUND_PATH,
    photo_cache_dir: Path | None = None,
) -> BinaryIO:
    presentation = Presentation()
    presentation.slide_width = Inches(SLIDE_WIDTH_INCHES)
    presentation.slide_height = Inches(SLIDE_HEIGHT_INCHES)
    blank_layout = presentation.slide_layouts[6]
    if students:
        cache_directory = photo_cache_dir
        if cache_directory is None and photo_loader is None:
            cache_directory = PHOTO_CACHE_DIR
        worker_count = min(PHOTO_DOWNLOAD_WORKERS, len(students))
        with TemporaryDirectory(prefix="bukang-pptx-photos-") as photo_directory:
            directory = Path(photo_directory)
            with ExitStack() as stack:
                cached_paths = [
                    cached_photo_path(student.photo_url, cache_directory)
                    if cache_directory is not None
                    else None
                    for student in students
                ]
                needs_download = [
                    bool(student.photo_url) and cached_paths[index] is None
                    for index, student in enumerate(students)
                ]
                client = (
                    stack.enter_context(create_photo_http_client())
                    if photo_loader is None and any(needs_download)
                    else None
                )
                executor = stack.enter_context(
                    ThreadPoolExecutor(max_workers=worker_count)
                )
                photo_futures = [
                    None
                    if not needs_download[index]
                    else executor.submit(
                        stream_photo_to_disk if photo_loader is None else spool_photo,
                        index,
                        student.photo_url,
                        directory,
                        client if photo_loader is None else photo_loader,
                    )
                    for index, student in enumerate(students)
                ]

                def resolved_photo(index: int) -> bytes | None:
                    cached_path = cached_paths[index]
                    if cached_path is not None:
                        try:
                            return cached_path.read_bytes()
                        except OSError:
                            return None

                    future = photo_futures[index]
                    raw_photo = consume_spooled_photo(
                        future.result() if future is not None else None
                    )
                    optimized = optimize_photo(raw_photo)
                    if optimized is not None and cache_directory is not None:
                        store_cached_photo(
                            students[index].photo_url,
                            optimized,
                            cache_directory,
                        )
                    return optimized

                for index in range(0, len(students), 2):
                    slide = presentation.slides.add_slide(blank_layout)
                    slide.shapes.add_picture(
                        str(background_path),
                        Inches(0),
                        Inches(0),
                        Inches(SLIDE_WIDTH_INCHES),
                        Inches(SLIDE_HEIGHT_INCHES),
                    )
                    first = students[index]
                    fill_slot(
                        slide,
                        SLOT_TOP,
                        _student_record(first),
                        resolved_photo(index),
                    )
                    if index + 1 < len(students):
                        second = students[index + 1]
                        fill_slot(
                            slide,
                            SLOT_BOTTOM,
                            _student_record(second),
                            resolved_photo(index + 1),
                        )

    output = TemporaryFile(mode="w+b")
    presentation.save(output)
    output.seek(0)
    return output
